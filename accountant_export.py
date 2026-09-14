import io
from datetime import datetime

import pandas as pd
from google.cloud import bigquery
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_ID = "gf-full-data"

VIEWS = {
    "All Transactions": f"{PROJECT_ID}.finance.accountant_transactions",
    "Monthly Summary": f"{PROJECT_ID}.finance.accountant_monthly_summary",
    "Annual Summary": f"{PROJECT_ID}.finance.accountant_annual_summary",
    "Source & Location": f"{PROJECT_ID}.finance.accountant_location_summary",
    "Gift Cards": f"{PROJECT_ID}.finance.accountant_gift_cards",
}


def query_view(client, view_name):
    query = f"""
        SELECT *
        FROM `{view_name}`
    """
    return client.query(query).to_dataframe()


def add_dataframe_sheet(wb, title, df):
    ws = wb.create_sheet(title=title)

    # Headers
    for col_num, column_name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_num, value=column_name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="202124")
        cell.alignment = Alignment(vertical="center")

    # Data
    for row_num, row in enumerate(df.itertuples(index=False, name=None), 2):
        for col_num, value in enumerate(row, 1):

            # Convert pandas NaN/NaT to blank Excel cells
            if pd.isna(value):
                value = None

            # Convert pandas timestamps to native Python datetime
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()

            ws.cell(
                row=row_num,
                column=col_num,
                value=value
            )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Table
    if len(df) > 0:
        end_col = get_column_letter(len(df.columns))
        end_row = len(df) + 1

        safe_table_name = (
            title
            .replace(" ", "")
            .replace("&", "And")
            .replace("-", "")
        )

        table = Table(
            displayName=f"{safe_table_name}Table",
            ref=f"A1:{end_col}{end_row}"
        )

        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        table.tableStyleInfo = style
        ws.add_table(table)

    # Formatting / widths
    for column_number, column_name in enumerate(df.columns, 1):
        col_letter = get_column_letter(column_number)

        lower = column_name.lower()

        width = 16

        if lower in {
            "transaction_id",
            "order_id",
        }:
            width = 28

        elif lower in {
            "location",
            "sales_location",
            "payment_method",
            "country",
            "order_status",
            "tax_method",
            "source_transaction_kind",
        }:
            width = 24

        elif lower in {
            "date",
            "month",
        }:
            width = 13

        elif "timestamp" in lower:
            width = 20

        ws.column_dimensions[col_letter].width = width

        # Date formatting
        if lower in {"date", "month"}:
            for cell in ws[col_letter][1:]:
                cell.number_format = "dd/mm/yyyy"

        elif "timestamp" in lower:
            for cell in ws[col_letter][1:]:
                cell.number_format = "dd/mm/yyyy hh:mm"

        # Financial formatting
        if any(
            key in lower
            for key in [
                "gross",
                "tax",
                "net",
                "discount",
                "shipping",
                "refund",
                "gift_card",
            ]
        ):
            for cell in ws[col_letter][1:]:
                cell.number_format = '#,##0.00;[Red]-#,##0.00'

    return ws


def add_notes_sheet(wb):
    ws = wb.create_sheet("Notes & Methodology", 0)

    ws.merge_cells("A1:F1")

    title = ws["A1"]
    title.value = "The Great Frog — Accountant Finance Export"
    title.font = Font(
        bold=True,
        color="FFFFFF",
        size=16
    )
    title.fill = PatternFill(
        "solid",
        fgColor="202124"
    )
    title.alignment = Alignment(
        vertical="center"
    )

    ws.row_dimensions[1].height = 28

    notes = [
        (
            "Purpose",
            "Unified historical transaction dataset for finance/accounting "
            "review, combining Shopify, Square and historical WooCommerce stores."
        ),
        (
            "Source of truth",
            "BigQuery finance.sales_master and accountant-facing views derived "
            "from it."
        ),
        (
            "Currencies",
            "GBP, USD and JPY are preserved in their original transaction "
            "currencies. No FX conversion to GBP is included."
        ),
        (
            "Sales / refunds",
            "Sales and refunds are separate transaction rows. Refund values "
            "are represented as negative amounts."
        ),
        (
            "Shopify",
            "Successful SALE and CAPTURE transactions are treated as money "
            "received. Successful REFUND transactions are money returned. "
            "AUTHORIZATION, FAILURE/ERROR and PENDING transactions are excluded."
        ),
        (
            "Migrated Shopify orders",
            "Matrixify-migrated WooCommerce orders are excluded from Shopify "
            "finance data to avoid double counting."
        ),
        (
            "Square VAT",
            "For standard-rated UK jewellery, VAT is derived from VAT-inclusive "
            "taxable sales where Square source tax was not reliable. Identified "
            "gift-card issuance is excluded from taxable/ex-VAT Square sales."
        ),
        (
            "Square gift cards",
            "Identified Square gift-card issuance is shown separately. This "
            "explains the expected difference between Square transaction gross "
            "and tax plus ex-tax sales."
        ),
        (
            "Shopify gift cards",
            "Shopify gift-card product sales are shown separately for "
            "information. Shopify tax/net values are not reduced again by this "
            "gift-card amount."
        ),
        (
            "WooCommerce gift cards",
            "Historical WooCommerce voucher/gift-card issuance is not separately "
            "identified in the current accountant views. The Gift Cards sheet "
            "therefore does not represent complete historical issuance."
        ),
        (
            "WooCommerce UK",
            "Historical UK/Worldwide WooCommerce sales preserve recorded source "
            "tax. Refund records were deduplicated before inclusion."
        ),
        (
            "WooCommerce US / JP",
            "USD and JPY remain in source currency. Recorded tax is retained "
            "where available."
        ),
        (
            "Migration pattern",
            "WooCommerce online sales transition to Shopify in November 2025, "
            "with trailing WooCommerce refunds afterward. POS migration from "
            "Square to Shopify is staggered through early 2026. Legitimate "
            "residual Square Wedding, Gold and legacy transactions remain."
        ),
        (
            "Tax treatment",
            "This workbook is a reporting dataset rather than a legal VAT "
            "determination. Gift-card/voucher treatment and unusual historical "
            "items should be confirmed by the accountant."
        ),
        (
            "Generated",
            datetime.now().strftime("%d/%m/%Y %H:%M")
        ),
    ]

    ws["A3"] = "Item"
    ws["B3"] = "Methodology / note"

    for cell in ws[3][:2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="202124")

    for row_num, (name, description) in enumerate(notes, 4):
        ws.cell(row_num, 1, name)
        ws.cell(row_num, 1).font = Font(bold=True)
        ws.cell(row_num, 1).fill = PatternFill(
            "solid",
            fgColor="E8EAED"
        )

        ws.cell(row_num, 2, description)
        ws.cell(row_num, 2).alignment = Alignment(
            wrap_text=True,
            vertical="top"
        )

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90

    ws["D3"] = "Workbook sheet"
    ws["E3"] = "What it contains"

    for cell in [ws["D3"], ws["E3"]]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="202124")

    sheet_notes = [
        (
            "All Transactions",
            "Full transaction-level ledger."
        ),
        (
            "Monthly Summary",
            "Monthly totals by source, channel, location and currency."
        ),
        (
            "Annual Summary",
            "Annual totals by source and currency."
        ),
        (
            "Source & Location",
            "Whole-period totals by source, channel, location and currency."
        ),
        (
            "Gift Cards",
            "Identified Shopify and Square gift-card issuance."
        ),
    ]

    for row_num, (sheet, description) in enumerate(sheet_notes, 4):
        ws.cell(row_num, 4, sheet)
        ws.cell(row_num, 4).font = Font(bold=True)
        ws.cell(row_num, 4).fill = PatternFill(
            "solid",
            fgColor="E8EAED"
        )

        ws.cell(row_num, 5, description)
        ws.cell(row_num, 5).alignment = Alignment(
            wrap_text=True,
            vertical="top"
        )

    ws.column_dimensions["D"].width = 25
    ws.column_dimensions["E"].width = 60

    ws.freeze_panes = "A4"


def build_accountant_export():
    client = bigquery.Client(project=PROJECT_ID)

    wb = Workbook()

    # Remove default empty sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    add_notes_sheet(wb)

    for sheet_name, view_name in VIEWS.items():
        print(f"Querying {view_name}")

        df = query_view(
            client,
            view_name
        )

        print(
            f"{sheet_name}: "
            f"{len(df):,} rows"
        )

        add_dataframe_sheet(
            wb,
            sheet_name,
            df
        )

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    return output
